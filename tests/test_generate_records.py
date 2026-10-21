import csv
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

import generate_records as gr

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"


class ReadTableTest(unittest.TestCase):
    def _csv(self, text):
        tmp = tempfile.NamedTemporaryFile("w", suffix=".csv", encoding="utf-8-sig", newline="", delete=False)
        tmp.write(text)
        tmp.close()
        return Path(tmp.name)

    def test_missing_column_raises(self):
        path = self._csv("이름,연락처\n홍,010\n")
        with self.assertRaises(ValueError) as ctx:
            gr.read_table(path, gr.MEMBER_COLUMNS, gr.MEMBER_REQUIRED, gr.MEMBER_DATES)
        self.assertIn("필요한 열이 없습니다", str(ctx.exception))

    def test_all_errors_reported_together(self):
        path = self._csv(
            "이름,생년월일,연락처,주소,가입일,상태\n"
            ",,010,,2024-02-01,정회원\n"
            "홍,1990-02-30,010,,2024-13-01,정회원\n"
        )
        with self.assertRaises(ValueError) as ctx:
            gr.read_table(path, gr.MEMBER_COLUMNS, gr.MEMBER_REQUIRED, gr.MEMBER_DATES)
        message = str(ctx.exception)
        self.assertIn("2행: '이름' 이(가) 비어 있습니다", message)
        self.assertIn("3행: '생년월일'", message)
        self.assertIn("3행: '가입일'", message)

    def test_blank_lines_are_skipped(self):
        path = self._csv("이름,생년월일,연락처,주소,가입일,상태\n,,,,,\n홍,,,,2024-01-01,\n")
        rows = gr.read_table(path, gr.MEMBER_COLUMNS, gr.MEMBER_REQUIRED, gr.MEMBER_DATES)
        self.assertEqual(len(rows), 1)


class SortingTest(unittest.TestCase):
    def test_members_sorted_by_join_date(self):
        members = gr.load_members(EXAMPLES / "members.csv")
        dates = [m["가입일"] for m in members]
        self.assertEqual(dates, sorted(dates))

    def test_history_sorted_by_date(self):
        history = gr.load_history(EXAMPLES / "history.csv")
        dates = [h["날짜"] for h in history]
        self.assertEqual(dates, sorted(dates))


class OutputTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_full_run_creates_three_documents(self):
        self.assertEqual(gr.main(["--data-dir", str(EXAMPLES), "--out-dir", str(self.out)]), 0)
        for name in ["회원명부.xlsx", "방문기록.docx", "연혁.docx"]:
            self.assertTrue((self.out / name).exists(), name)

    def test_member_workbook_contents(self):
        members = gr.load_members(EXAMPLES / "members.csv")
        path = self.out / "회원명부.xlsx"
        gr.write_members_xlsx(members, path, "테스트")

        wb = load_workbook(path)
        ws = wb["회원명부"]
        self.assertEqual([c.value for c in ws[1]], ["번호"] + gr.MEMBER_COLUMNS)
        self.assertEqual(ws.max_row, len(members) + 1)
        self.assertEqual(ws["A2"].value, 1)
        self.assertEqual(ws["B2"].value, members[0]["이름"])
        self.assertEqual(ws.freeze_panes, "A2")

        summary = wb["요약"]
        self.assertEqual(summary["B2"].value, len(members))

    def test_visit_table_rows(self):
        visits = gr.load_visits(EXAMPLES / "visits.csv")
        path = self.out / "방문기록.docx"
        gr.write_visits_docx(visits, path, "테스트")

        table = Document(path).tables[0]
        self.assertEqual(len(table.rows), len(visits) + 1)
        self.assertEqual(table.rows[0].cells[0].text, "번호")

    def test_history_grouped_by_year(self):
        history = gr.load_history(EXAMPLES / "history.csv")
        path = self.out / "연혁.docx"
        gr.write_history_docx(history, path, "테스트")

        headings = [p.text for p in Document(path).paragraphs if p.style.name == "Heading 2"]
        self.assertEqual(headings, ["2024년", "2025년", "2026년"])

    def test_empty_visits_writes_notice(self):
        path = self.out / "빈방문.docx"
        gr.write_visits_docx([], path, "테스트")
        texts = [p.text for p in Document(path).paragraphs]
        self.assertIn("기록이 없습니다.", texts)

    def test_missing_data_dir_fails(self):
        with self.assertRaises(SystemExit):
            gr.main(["--data-dir", str(self.out / "없는폴더"), "--out-dir", str(self.out)])


if __name__ == "__main__":
    unittest.main()
