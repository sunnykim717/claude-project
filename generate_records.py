#!/usr/bin/env python3
"""회원명부(엑셀), 방문기록과 연혁(워드) 문서를 CSV 데이터로부터 생성합니다.

사용법:
    python generate_records.py                       # private/ 에서 읽어 out/ 에 저장
    python generate_records.py --data-dir examples   # 예시 데이터로 시험

개인정보가 들어 있는 CSV 는 private/ 에 두세요. private/ 와 out/ 은 git 에 올라가지 않습니다.
"""

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KST = timezone(timedelta(hours=9))

MEMBER_COLUMNS = ["이름", "생년월일", "연락처", "주소", "가입일", "상태"]
MEMBER_REQUIRED = ["이름", "가입일"]
MEMBER_DATES = ["생년월일", "가입일"]

VISIT_COLUMNS = ["날짜", "시간", "방문자", "소속", "방문목적", "담당자"]
VISIT_REQUIRED = ["날짜", "방문자"]
VISIT_DATES = ["날짜"]

HISTORY_COLUMNS = ["날짜", "구분", "내용"]
HISTORY_REQUIRED = ["날짜", "내용"]
HISTORY_DATES = ["날짜"]


def is_valid_date(value):
    """YYYY-MM-DD 형식이고 실제로 존재하는 날짜인지 확인합니다."""
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", value):
        return False
    try:
        datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        return False
    return True


def read_table(path, columns, required, date_fields):
    """CSV 를 읽어 열 구조와 값을 검사합니다. 문제는 모아서 한 번에 알려 줍니다."""
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        missing = [c for c in columns if c not in (reader.fieldnames or [])]
        if missing:
            raise ValueError(f"{path.name}: 필요한 열이 없습니다 -> {', '.join(missing)}")

        rows, errors = [], []
        for line_no, raw in enumerate(reader, start=2):
            row = {c: (raw.get(c) or "").strip() for c in columns}
            if not any(row.values()):
                continue  # 빈 줄은 건너뜁니다.
            for field in required:
                if not row[field]:
                    errors.append(f"{path.name} {line_no}행: '{field}' 이(가) 비어 있습니다.")
            for field in date_fields:
                if row[field] and not is_valid_date(row[field]):
                    errors.append(f"{path.name} {line_no}행: '{field}' 은(는) YYYY-MM-DD 형식이어야 합니다 ({row[field]}).")
            rows.append(row)

    if errors:
        raise ValueError("\n".join(errors))
    return rows


def load_members(path):
    rows = read_table(path, MEMBER_COLUMNS, MEMBER_REQUIRED, MEMBER_DATES)
    return sorted(rows, key=lambda r: (r["가입일"], r["이름"]))


def load_visits(path):
    rows = read_table(path, VISIT_COLUMNS, VISIT_REQUIRED, VISIT_DATES)
    return sorted(rows, key=lambda r: (r["날짜"], r["시간"]))


def load_history(path):
    rows = read_table(path, HISTORY_COLUMNS, HISTORY_REQUIRED, HISTORY_DATES)
    return sorted(rows, key=lambda r: r["날짜"])


# ---------- 엑셀: 회원명부 ----------

def write_members_xlsx(members, path, generated_at):
    wb = Workbook()
    ws = wb.active
    ws.title = "회원명부"

    headers = ["번호"] + MEMBER_COLUMNS
    ws.append(headers)
    for number, member in enumerate(members, start=1):
        ws.append([number] + [member[c] for c in MEMBER_COLUMNS])

    header_fill = PatternFill("solid", fgColor="DBEAFE")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, header in enumerate(headers, start=1):
        longest = max([len(header)] + [len(str(ws.cell(row=r, column=index).value or "")) for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(index)].width = min(max(longest * 1.6 + 2, 8), 50)

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("요약")
    summary.append(["항목", "값"])
    summary.append(["총 회원 수", len(members)])
    summary.append(["작성일", generated_at])
    counts = defaultdict(int)
    for member in members:
        counts[member["상태"] or "(미지정)"] += 1
    for status, count in sorted(counts.items()):
        summary.append([f"상태: {status}", count])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 22

    wb.save(path)


# ---------- 워드: 방문기록, 연혁 ----------

def write_visits_docx(visits, path, generated_at):
    doc = Document()
    doc.add_heading("방문 기록", level=1)
    doc.add_paragraph(f"작성일: {generated_at}    총 {len(visits)}건")

    if not visits:
        doc.add_paragraph("기록이 없습니다.")
    else:
        headers = ["번호"] + VISIT_COLUMNS
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for number, visit in enumerate(visits, start=1):
            cells = table.add_row().cells
            for cell, value in zip(cells, [str(number)] + [visit[c] for c in VISIT_COLUMNS]):
                cell.text = value

    doc.save(path)


def write_history_docx(history, path, generated_at):
    doc = Document()
    doc.add_heading("연혁", level=1)
    doc.add_paragraph(f"작성일: {generated_at}")

    if not history:
        doc.add_paragraph("기록이 없습니다.")
    else:
        by_year = defaultdict(list)
        for item in history:
            by_year[item["날짜"][:4]].append(item)
        for year in sorted(by_year):
            doc.add_heading(f"{year}년", level=2)
            for item in by_year[year]:
                kind = f"[{item['구분']}] " if item["구분"] else ""
                doc.add_paragraph(f"{item['날짜']}  {kind}{item['내용']}", style="List Bullet")

    doc.save(path)


# ---------- 실행 ----------

def main(argv=None):
    parser = argparse.ArgumentParser(description="회원명부, 방문기록, 연혁 문서 자동 생성")
    parser.add_argument("--data-dir", type=Path, default=Path("private"), help="CSV 가 있는 폴더")
    parser.add_argument("--out-dir", type=Path, default=Path("out"), help="문서를 저장할 폴더")
    args = parser.parse_args(argv)

    generated_at = datetime.now(KST).strftime("%Y-%m-%d %H:%M (KST)")
    args.out_dir.mkdir(parents=True, exist_ok=True)

    jobs = [
        ("members.csv", load_members, write_members_xlsx, "회원명부.xlsx"),
        ("visits.csv", load_visits, write_visits_docx, "방문기록.docx"),
        ("history.csv", load_history, write_history_docx, "연혁.docx"),
    ]
    produced = 0
    for source, load, write, output_name in jobs:
        source_path = args.data_dir / source
        if not source_path.exists():
            print(f"건너뜀: {source_path} 파일이 없습니다.")
            continue
        records = load(source_path)
        write(records, args.out_dir / output_name, generated_at)
        print(f"생성 완료: {args.out_dir / output_name} ({len(records)}건)")
        produced += 1

    if produced == 0:
        raise SystemExit(f"{args.data_dir} 에서 처리할 CSV 파일을 찾지 못했습니다.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
